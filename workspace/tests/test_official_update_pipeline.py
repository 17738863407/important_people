from __future__ import annotations

import importlib.util
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parents[2]


def load_script(filename: str, module_name: str):
    script_path = PROJECT_DIR / filename
    spec = importlib.util.spec_from_file_location(module_name, script_path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_upload_mapping_and_in_memory_workbook() -> None:
    module = load_script(
        "script4-职务比对1-上传官员信息.py", "official_update_upload"
    )
    header, rows = module.read_tsv_rows(PROJECT_DIR / "职务规范4：扩充.txt")

    assert header == module.EXPECTED_SOURCE_HEADER
    upload_header, upload_rows = module.select_upload_data(header, rows[:2])
    assert upload_header == [
        "姓名",
        "职务",
        "依据说明",
        "依据链接",
        "排序",
        "所属省",
        "所属市",
    ]

    workbook_bytes = module.generate_excel_bytes(upload_header, upload_rows)
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True)
    try:
        worksheet = workbook["官员信息"]
        values = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    assert list(values[0]) == upload_header
    assert len(values) == len(upload_rows) + 1


def test_incomplete_online_fetch_does_not_overwrite(tmp_path, monkeypatch) -> None:
    module = load_script(
        "script4-职务比对2-获取在线比对数据.py", "official_update_fetch"
    )
    output_path = tmp_path / "existing.txt"
    output_path.write_text("KEEP", encoding="utf-8")
    token_file = tmp_path / ".secrets" / "official_update_token"
    token_file.parent.mkdir()
    token_file.write_text("test-token", encoding="utf-8")

    class FakeSession:
        trust_env = False

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

    monkeypatch.setattr(module.requests, "Session", FakeSession)
    monkeypatch.setattr(module, "fetch_all", lambda *_args: ([{"name": "甲"}], 2, [2]))

    try:
        module.main(
            "http://example.invalid",
            20,
            1,
            output_path,
            tmp_path / "log",
            token_file,
            tmp_path / "chrome.exe",
            tmp_path / "profile",
            9777,
            "http://example.invalid/page",
            "example.invalid",
            1,
            1,
            False,
        )
    except RuntimeError as error:
        assert "抓取不完整" in str(error)
    else:
        raise AssertionError("分页不完整时应拒绝写出结果")

    assert output_path.read_text(encoding="utf-8") == "KEEP"


def test_filter_and_normalize_pipeline(tmp_path) -> None:
    filter_module = load_script(
        "script4-职务比对3-获取无重复数据.py", "official_update_filter"
    )
    normalize_module = load_script(
        "script4-职务比对4-无重复数据规范化.py", "official_update_normalize"
    )

    online_path = tmp_path / "online.txt"
    filtered_path = tmp_path / "filtered.txt"
    filter_log = tmp_path / "filter.log"
    online_path.write_text(
        "\t".join(filter_module.__dict__.get("SPLIT_HEADERS", [
            "姓名", "本地职务", "职级", "排序", "所属省", "所属市", "提醒类型",
            "老数据职务", "依据链接", "需要更新时间", "操作"
        ]))
        + "\n"
        + "甲\t新职务\t正厅级\t1\t北京\t\t需要新增\t旧职务\thttps://a.example\t\t删除\n"
        + "乙\t相同职务\t副厅级\t2\t北京\t\t职务变更\t相同职务\thttps://b.example\t\t删除\n",
        encoding="utf-8",
    )
    filter_module.main(online_path, filtered_path, filter_log)
    filtered_lines = filtered_path.read_text(encoding="utf-8").splitlines()
    assert len(filtered_lines) == 2
    assert filtered_lines[1].startswith("甲\t新职务\t")

    normalized_path = tmp_path / "normalized.txt"
    normalized_path.write_text(
        "姓名*\t职务*\t排序*\t职级*\t所属省*\t所属市\t依据链接\t依据说明\t维护人\t机构\t详情页链接\t脚本序号\n"
        "甲\t新职务\t1\t正厅级\t北京\t\thttps://a.example\t网页信息\t维护员\t北京市人民政府\t\t0001\n",
        encoding="utf-8",
    )
    output_xlsx = tmp_path / "normalized.xlsx"
    normalize_module.main(
        filtered_path,
        normalized_path,
        output_xlsx,
        tmp_path / "normalize.log",
    )

    workbook = load_workbook(output_xlsx, read_only=True)
    try:
        values = list(workbook["无重复排序"].iter_rows(values_only=True))
    finally:
        workbook.close()
    assert len(values) == 2
    assert values[0][0:3] == ("姓名*", "职务*", "老数据职务")
    assert values[1][0:3] == ("甲", "新职务", "旧职务")
